"""
Build the print-ready A4 math worksheet (.docx).

Layout overview:
- A4 portrait, 15 mm margins (good for home/office printers).
- A small header (Name / Date / Score) followed by a centred title.
- Problems are placed in an outer grid (3 columns by default). Each cell of the
  outer grid contains a "problem card": a problem number plus a nested table
  drawn as the working area:

      ┌──┬──┬──┬──┐  ← thin carry row (dashed light borders)
      │  │  │  │  │
      ├──┼──┼──┼──┤
      │  │  │ 9│ 8│  ← operand 1 (no borders, digits right-aligned)
      │ +│  │ 7│ 8│  ← operand 2, operator in leftmost cell
      ├──┼──┼──┼──┤  ← solid rule
      │  │  │  │  │  ← answer boxes (solid borders, large)
      └──┴──┴──┴──┘

- For ÷ we use a horizontal format with an answer box: "12 ÷ 3 =  ⬜".
- Black-and-white only. All output is print-ready as-is.
- A separate "Answer Key" page is appended when requested.

Implementation notes:
- python-docx does not expose cell borders directly, so we set them via OXML.
- All widths are in millimetres for predictability on A4.
"""

from __future__ import annotations

from io import BytesIO
from typing import Iterable

from docx import Document
from docx.document import Document as _Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.table import _Cell, Table

from .generator import GenerationRequest, Problem, generate, max_digits

# ---- layout constants (millimetres) -----------------------------------------
PAGE_MARGIN_MM = 15
PROBLEMS_PER_ROW = 3
DIGIT_CELL_MM = 9            # box for one digit in operand/answer rows
DIGIT_ROW_HEIGHT_MM = 10     # row height for operand & answer rows (a bit taller)
CARRY_ROW_HEIGHT_MM = 5      # thin carry row
OUTER_CELL_MARGIN_TOP = 200      # DXA — vertical breathing room around each problem card
OUTER_CELL_MARGIN_BOTTOM = 280   # DXA — extra below, so the answer boxes don't sit against the next row
OUTER_CELL_MARGIN_SIDE = 140     # DXA — horizontal gutter between problems
PROBLEM_CARD_TRAILING_PT = 10    # empty paragraph after each card for safety spacing

# ---- low-level XML helpers --------------------------------------------------

def _qn(tag: str) -> str:
    return qn(tag)


# OOXML schema-mandated child order for <w:tcPr> and <w:tcBorders>.
# Word/LibreOffice opens out-of-order files but they fail strict validation.
_TC_PR_ORDER = [
    "cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
    "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark",
    "headers",
]
_TC_BORDERS_ORDER = [
    "top", "start", "left", "bottom", "end", "right", "insideH", "insideV",
    "tl2br", "tr2bl",
]


def _local_name(el) -> str:
    """Return the tag's local name, stripping any namespace."""
    tag = el.tag
    return tag.split("}", 1)[1] if "}" in tag else tag


def _insert_in_order(parent, child, schema_order: list[str]) -> None:
    """Insert child into parent at the position required by schema_order."""
    name = _local_name(child)
    idx = schema_order.index(name)
    # remove any existing element with the same local name
    for existing in list(parent):
        if _local_name(existing) == name:
            parent.remove(existing)
    # find the first sibling whose schema position is after ours
    for sibling in list(parent):
        sib_name = _local_name(sibling)
        if sib_name in schema_order and schema_order.index(sib_name) > idx:
            sibling.addprevious(child)
            return
    parent.append(child)


def _set_cell_border(cell: _Cell, **edges) -> None:
    """
    Set borders on a single cell.

    Each kwarg is one of top / left / bottom / right, mapped to a dict like:
        {"val": "single", "sz": "8", "color": "000000"}
    `val` may be "single", "dashed", "dotted", or "nil" (no border).
    `sz` is in eighths of a point (8 == 1pt).
    """
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge_name in _TC_BORDERS_ORDER:
        if edge_name not in edges:
            continue
        edge_el = OxmlElement(f"w:{edge_name}")
        for k, v in edges[edge_name].items():
            edge_el.set(_qn(f"w:{k}"), str(v))
        tc_borders.append(edge_el)
    _insert_in_order(tc_pr, tc_borders, _TC_PR_ORDER)


def _set_cell_margins(cell: _Cell, top=20, bottom=20, left=40, right=40) -> None:
    """Set cell internal margins (in DXA: 1440 = 1 inch). Small but readable."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    # children of tcMar must appear in this order: top, start, left, bottom, end, right
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(_qn("w:w"), str(val))
        el.set(_qn("w:type"), "dxa")
        tc_mar.append(el)
    _insert_in_order(tc_pr, tc_mar, _TC_PR_ORDER)


def _no_border() -> dict:
    return {"val": "nil"}


def _solid(sz: int = 8, color: str = "000000") -> dict:
    return {"val": "single", "sz": str(sz), "color": color}


def _dashed_light() -> dict:
    return {"val": "dashed", "sz": "4", "color": "808080"}


# ---- page setup -------------------------------------------------------------

def _setup_page(doc: _Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)   # A4 portrait
    section.page_height = Mm(297)
    section.top_margin = Mm(PAGE_MARGIN_MM)
    section.bottom_margin = Mm(PAGE_MARGIN_MM)
    section.left_margin = Mm(PAGE_MARGIN_MM)
    section.right_margin = Mm(PAGE_MARGIN_MM)

    # default style: black text on white, readable size
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor(0x00, 0x00, 0x00)


# ---- header section ---------------------------------------------------------

def _add_header_block(doc: _Document, title: str) -> None:
    # Name / Date / Score line, then the title centred under it.
    line = doc.add_paragraph()
    line.paragraph_format.space_after = Pt(2)
    run = line.add_run("Name: ____________________________"
                       "    Date: __________________"
                       "    Score: ______")
    run.font.size = Pt(11)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(4)
    title_p.paragraph_format.space_after = Pt(6)
    title_run = title_p.add_run(title)
    title_run.bold = True
    title_run.font.size = Pt(18)


def _add_section_heading(doc: _Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)


# ---- per-problem card -------------------------------------------------------

def _set_row_height(row, height_mm: float, exact: bool = True) -> None:
    row.height = Mm(height_mm)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY if exact else WD_ROW_HEIGHT_RULE.AT_LEAST


def _digit_runs(cell: _Cell, text: str, size_pt: int = 18, bold: bool = False) -> None:
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.bold = bold


def _build_vertical_problem(parent_cell: _Cell, p: Problem, number: int, digit_cols: int) -> None:
    """
    Render a +, -, or × problem inside parent_cell using the boxed-digit layout.
    digit_cols is the global max number of digit columns (so all cards line up).
    """
    # problem number line
    num_p = parent_cell.add_paragraph()
    num_p.paragraph_format.space_before = Pt(0)
    num_p.paragraph_format.space_after = Pt(2)
    num_run = num_p.add_run(f"{number}.")
    num_run.bold = True
    num_run.font.size = Pt(11)

    # nested table: (operator col) + digit_cols columns × 4 rows
    total_cols = digit_cols + 1
    inner = parent_cell.add_table(rows=4, cols=total_cols)
    inner.alignment = WD_TABLE_ALIGNMENT.LEFT
    inner.autofit = False

    for col in inner.columns:
        for c in col.cells:
            c.width = Mm(DIGIT_CELL_MM)

    carry_row, op1_row, op2_row, ans_row = inner.rows
    _set_row_height(carry_row, CARRY_ROW_HEIGHT_MM)
    _set_row_height(op1_row, DIGIT_ROW_HEIGHT_MM)
    _set_row_height(op2_row, DIGIT_ROW_HEIGHT_MM)
    _set_row_height(ans_row, DIGIT_ROW_HEIGHT_MM)

    # right-align digit strings into the digit columns (skip col 0 = operator col)
    op1_digits = list(str(p.operand1))
    op2_digits = list(str(p.operand2))
    pad1 = [""] * (digit_cols - len(op1_digits))
    pad2 = [""] * (digit_cols - len(op2_digits))
    op1_cells = pad1 + op1_digits
    op2_cells = pad2 + op2_digits

    # ---- carry row: small, light dashed borders on digit cells only ----
    for ci, cell in enumerate(carry_row.cells):
        _set_cell_margins(cell, top=10, bottom=10, left=20, right=20)
        if ci == 0:
            _set_cell_border(cell,
                             top=_no_border(), left=_no_border(),
                             bottom=_no_border(), right=_no_border())
        else:
            _set_cell_border(cell,
                             top=_dashed_light(), left=_dashed_light(),
                             bottom=_dashed_light(), right=_dashed_light())
        _digit_runs(cell, "", size_pt=10)

    # ---- operand 1 row: no borders, digits right-aligned in cells ----
    for ci, cell in enumerate(op1_row.cells):
        _set_cell_margins(cell)
        _set_cell_border(cell,
                         top=_no_border(), left=_no_border(),
                         bottom=_no_border(), right=_no_border())
        if ci == 0:
            _digit_runs(cell, "", size_pt=18)
        else:
            _digit_runs(cell, op1_cells[ci - 1], size_pt=18)

    # ---- operand 2 row: operator in col 0, digits in remaining cols ----
    #     The bottom border on this whole row provides the horizontal rule.
    for ci, cell in enumerate(op2_row.cells):
        _set_cell_margins(cell)
        _set_cell_border(cell,
                         top=_no_border(), left=_no_border(),
                         bottom=_solid(sz=12), right=_no_border())
        if ci == 0:
            _digit_runs(cell, p.operator, size_pt=18, bold=True)
        else:
            _digit_runs(cell, op2_cells[ci - 1], size_pt=18)

    # ---- answer row: full bordered boxes (skip col 0 to keep operator col blank) ----
    for ci, cell in enumerate(ans_row.cells):
        _set_cell_margins(cell)
        if ci == 0:
            _set_cell_border(cell,
                             top=_no_border(), left=_no_border(),
                             bottom=_no_border(), right=_no_border())
            _digit_runs(cell, "", size_pt=18)
        else:
            _set_cell_border(cell,
                             top=_solid(sz=8), left=_solid(sz=8),
                             bottom=_solid(sz=8), right=_solid(sz=8))
            _digit_runs(cell, "", size_pt=18)


def _build_division_problem(parent_cell: _Cell, p: Problem, number: int) -> None:
    """Simple horizontal division: '12 ÷ 3 =  ⬜' with a bordered answer box."""
    num_p = parent_cell.add_paragraph()
    num_p.paragraph_format.space_before = Pt(0)
    num_p.paragraph_format.space_after = Pt(2)
    num_run = num_p.add_run(f"{number}.")
    num_run.bold = True
    num_run.font.size = Pt(11)

    # tiny table: [ "12 ÷ 3 = " | answer box ]
    inner = parent_cell.add_table(rows=1, cols=2)
    inner.alignment = WD_TABLE_ALIGNMENT.LEFT
    inner.autofit = False

    # text cell
    text_cell = inner.cell(0, 0)
    text_cell.width = Mm(DIGIT_CELL_MM * 3)
    _set_cell_margins(text_cell)
    _set_cell_border(text_cell, top=_no_border(), left=_no_border(),
                     bottom=_no_border(), right=_no_border())
    _digit_runs(text_cell, f"{p.operand1} ÷ {p.operand2} =", size_pt=16)
    text_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    # answer box cell
    box_cell = inner.cell(0, 1)
    box_cell.width = Mm(DIGIT_CELL_MM * max(2, len(str(p.answer)) + 1))
    _set_row_height(inner.rows[0], DIGIT_ROW_HEIGHT_MM)
    _set_cell_margins(box_cell)
    _set_cell_border(box_cell, top=_solid(sz=8), left=_solid(sz=8),
                     bottom=_solid(sz=8), right=_solid(sz=8))
    _digit_runs(box_cell, "", size_pt=18)


def _add_problem_grid(doc: _Document, problems: list[Problem], digit_cols: int,
                      start_number: int) -> int:
    """Place a list of problems into a PROBLEMS_PER_ROW-wide outer table."""
    if not problems:
        return start_number

    rows_needed = (len(problems) + PROBLEMS_PER_ROW - 1) // PROBLEMS_PER_ROW
    outer = doc.add_table(rows=rows_needed, cols=PROBLEMS_PER_ROW)
    outer.alignment = WD_TABLE_ALIGNMENT.LEFT
    outer.autofit = False

    # available content width in mm
    content_mm = 210 - 2 * PAGE_MARGIN_MM
    cell_mm = content_mm / PROBLEMS_PER_ROW
    for col in outer.columns:
        for c in col.cells:
            c.width = Mm(cell_mm)

    num = start_number
    for i, p in enumerate(problems):
        r = i // PROBLEMS_PER_ROW
        c = i % PROBLEMS_PER_ROW
        cell = outer.cell(r, c)
        _set_cell_margins(cell,
                          top=OUTER_CELL_MARGIN_TOP,
                          bottom=OUTER_CELL_MARGIN_BOTTOM,
                          left=OUTER_CELL_MARGIN_SIDE,
                          right=OUTER_CELL_MARGIN_SIDE)
        # clear default paragraph so our content starts cleanly
        cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
        _set_cell_border(cell,
                         top=_no_border(), left=_no_border(),
                         bottom=_no_border(), right=_no_border())
        if p.operator == "÷":
            _build_division_problem(cell, p, num)
        else:
            _build_vertical_problem(cell, p, num, digit_cols)
        # trailing empty paragraph inside the card for extra vertical breathing room
        trailing = cell.add_paragraph()
        trailing.paragraph_format.space_before = Pt(0)
        trailing.paragraph_format.space_after = Pt(PROBLEM_CARD_TRAILING_PT)
        num += 1
    return num


# ---- answer key -------------------------------------------------------------

def _add_answer_key(doc: _Document, problems: list[Problem]) -> None:
    # page break
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("Answer Key")
    title_run.bold = True
    title_run.font.size = Pt(18)

    # multi-column list of answers
    cols = 4
    rows = (len(problems) + cols - 1) // cols
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    content_mm = 210 - 2 * PAGE_MARGIN_MM
    cell_mm = content_mm / cols
    for col in table.columns:
        for c in col.cells:
            c.width = Mm(cell_mm)
            _set_cell_border(c, top=_no_border(), left=_no_border(),
                             bottom=_no_border(), right=_no_border())

    for i, prob in enumerate(problems):
        r = i // cols
        c = i % cols
        cell = table.cell(r, c)
        cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
        line = cell.add_paragraph()
        line.paragraph_format.space_after = Pt(2)
        run = line.add_run(
            f"{i + 1}.  {prob.operand1} {prob.operator} {prob.operand2} = {prob.answer}"
        )
        run.font.size = Pt(11)


# ---- top-level ---------------------------------------------------------------

def build_docx(req: GenerationRequest) -> bytes:
    problems = generate(req)
    if not problems:
        raise ValueError("No problems were generated; check the operation settings.")

    # group by operator, preserving the order add → sub → mul → div
    grouping_order = ["+", "-", "×", "÷"]
    groups: dict[str, list[Problem]] = {op: [] for op in grouping_order}
    for prob in problems:
        groups[prob.operator].append(prob)
    for op in groups:
        groups[op].sort(key=lambda p: p.difficulty)

    # shared digit-column width across all vertical problems for visual alignment
    vertical_problems = [p for p in problems if p.operator != "÷"]
    digit_cols = max_digits(vertical_problems) if vertical_problems else 1

    doc = Document()
    _setup_page(doc)
    _add_header_block(doc, req.title)

    section_titles = {
        "+": "Addition",
        "-": "Subtraction",
        "×": "Multiplication",
        "÷": "Division",
    }

    # the worksheet is rendered in this exact order; the answer key must follow it
    ordered_problems: list[Problem] = []
    number = 1
    first_section = True
    for op in grouping_order:
        section_problems = groups[op]
        if not section_problems:
            continue
        if not first_section:
            doc.add_paragraph()  # gap between sections
        _add_section_heading(doc, section_titles[op])
        number = _add_problem_grid(doc, section_problems, digit_cols, number)
        ordered_problems.extend(section_problems)
        first_section = False

    if req.include_answer_key:
        _add_answer_key(doc, ordered_problems)

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
